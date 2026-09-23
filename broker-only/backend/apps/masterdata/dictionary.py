"""Load the Unit Attribute Dictionary from the founder-review spreadsheet or the committed YAML."""

from pathlib import Path

import yaml

from .models import AttributeDef

APPLIES = {"R": "RENT", "SN": "SALE_NEW", "SR": "SALE_RESALE"}
DEFAULT_PATH = Path(__file__).parent / "dictionary" / "attribute_dictionary.yaml"


def rows_from_xlsx(path, *, approved_only: bool) -> list[dict]:
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True, data_only=True)["Attributes"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip() if h else "" for h in rows[0]]
    idx = {h: i for i, h in enumerate(hdr)}
    out = []
    for r in rows[1:]:
        key = r[idx["Key"]]
        if not key:
            continue
        decision = (r[idx["Your decision"]] or "").strip()
        if approved_only and decision not in ("Approve", "Modify"):
            continue
        vtype = r[idx["Type"]]
        values = r[idx["Allowed values / unit"]] or ""
        is_choice = vtype in ("enum", "multi_enum")
        out.append(
            {
                "key": key,
                "label": r[idx["Label"]],
                "category": r[idx["Category"]],
                "scope": r[idx["Scope"]],
                "value_type": vtype,
                "allowed_values": [v.strip() for v in values.split("|")] if is_choice and values else [],
                "unit_label": "" if is_choice else values,
                "authority": r[idx["Authority"]],
                "matching": (r[idx["Matching"]] or "Display").lower(),
                "public": r[idx["Public"]] or "N",
                "applies_to": [APPLIES[a.strip()] for a in (r[idx["Applies to"]] or "").split(",") if a.strip()],
                "mvp": (r[idx["MVP"]] or "N") == "Y",
                "notes": r[idx["Notes"]] or "",
                "review": {"decision": decision, "comments": r[idx["Your comments"]] or ""},
            }
        )
    return out


def derive_entry(row: dict) -> tuple[str, str]:
    """Tier and who-is-asked, derived from the approved columns (docs/06 rule 3 and 4)."""
    if row["authority"] == "computed":
        return "system", "nobody"
    if row["mvp"] and row["matching"] == "hard":
        tier = "essential"
    elif row["mvp"]:
        tier = "recommended"
    else:
        tier = "detailed"
    if row["scope"] in ("building", "society"):
        asked = "building"
    elif row["authority"] == "owner" and "house rule" in row["category"].lower():
        asked = "owner"
    else:
        asked = "broker"
    return tier, asked


def load_rows(rows: list[dict], version: str) -> tuple[int, int]:
    created = updated = 0
    keys = set()
    for order, row in enumerate(rows):
        row = dict(row)
        row.pop("review", None)
        tier, asked = derive_entry(row)
        row.setdefault("entry_tier", tier)
        row.setdefault("asked_of", asked)
        row.setdefault("display_order", order)
        keys.add(row["key"])
        _, was_created = AttributeDef.objects.update_or_create(
            key=row.pop("key"), defaults={**row, "active": True, "dictionary_version": version}
        )
        created += was_created
        updated += not was_created
    # Attributes dropped from the dictionary are deactivated, never deleted (observations reference them).
    AttributeDef.objects.exclude(key__in=keys).update(active=False)
    return created, updated


def rows_from_yaml(path) -> tuple[list[dict], str]:
    doc = yaml.safe_load(Path(path).read_text(encoding="utf-8"))
    return doc["attributes"], str(doc.get("version", ""))


def write_yaml(rows: list[dict], path, version: str) -> None:
    clean = [{k: v for k, v in r.items() if k != "review"} for r in rows]
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(yaml.safe_dump({"version": version, "attributes": clean}, allow_unicode=True, sort_keys=False), encoding="utf-8")
