from decimal import Decimal, InvalidOperation
from django.db import transaction

from openpyxl import load_workbook

from society.models import (
    Flat,
    FlatOwnership,
    FlatOwner,
)


EXPECTED_HEADERS = [
    "society id",
    "wing",
    "flat no",
    "owner name",
    "phone",
    "maintenance outstanding",
    "interest outstanding",
    "penalty outstanding",
    "advance balance",
]


def normalize_header(value):

    if value is None:
        return ""

    return str(value).strip().lower()


def normalize_amount(value):

    if value is None:
        return Decimal("0.00")

    value = str(value).strip()

    if value == "":
        return Decimal("0.00")

    value = value.replace(",", "")

    try:
        amount = Decimal(value)
    except InvalidOperation:
        raise ValueError("Invalid numeric value")

    if amount < 0:
        raise ValueError("Negative values are not allowed")

    return amount


def parse_maintenance_bill_balance_excel(file_obj, society):

    try:
        workbook = load_workbook(file_obj)

        worksheet = workbook.active

    except Exception:

        return {
            "status": "failed",
            "errors": ["Invalid Excel file"],
            "warnings": [],
            "rows": [],
        }

    rows = list(
        worksheet.iter_rows(values_only=True)
    )

    if not rows:

        return {
            "status": "failed",
            "errors": ["Excel file is empty"],
            "warnings": [],
            "rows": [],
        }

    raw_headers = rows[0]

    normalized_headers = [
        normalize_header(header)
        for header in raw_headers
    ]

    if normalized_headers != EXPECTED_HEADERS:

        return {
            "status": "failed",
            "errors": [
                "Excel headers do not match expected format"
            ],
            "warnings": [],
            "rows": [],
        }

    parsed_rows = []

    errors = []

    warnings = []

    processed_flats = set()

    for index, row in enumerate(rows[1:], start=2):

        if all(
            cell is None or str(cell).strip() == ""
            for cell in row
        ):
            continue

        try:

            society_id = row[0]
            wing = row[1]
            flat_no_raw = row[2]
            owner_name_raw = row[3]
            phone_raw = row[4]

            maintenance_raw = row[5]
            interest_raw = row[6]
            penalty_raw = row[7]
            advance_raw = row[8]

            flat_no = (
                str(flat_no_raw).strip()
                if flat_no_raw else ""
            )

            owner_name = (
                str(owner_name_raw).strip()
                if owner_name_raw else ""
            )

            phone = (
                str(phone_raw).strip()
                if phone_raw else ""
            )

            if not flat_no:

                errors.append(
                    f"Row {index}: Flat No missing"
                )

                continue

            if flat_no in processed_flats:

                errors.append(
                    f"Row {index}: Duplicate flat in upload ({flat_no})"
                )

                continue

            processed_flats.add(flat_no)

            flat = Flat.objects.filter(
                society=society,
                flat_number=flat_no,
            ).first()

            if not flat:

                errors.append(
                    f"Row {index}: Flat not found ({flat_no})"
                )

                continue

            ownership = (
                FlatOwnership.objects.filter(
                    flat=flat,
                    is_active=True,
                ).first()
            )

            if not ownership:

                errors.append(
                    f"Row {index}: No active ownership found ({flat_no})"
                )

                continue

            system_owner = (
                FlatOwner.objects.filter(
                    ownership=ownership
                )
                .select_related("person")
                .first()
            )

            if system_owner:

                db_owner_name = ""

                db_phone = ""

                if system_owner.person:
                    db_owner_name = (
                        system_owner.person.full_name
                    )

                    db_phone = (
                        system_owner.person.phone
                    )

                elif system_owner.legal_entity_name:
                    db_owner_name = (
                        system_owner.legal_entity_name
                    )

                if (
                    owner_name
                    and owner_name.lower()
                    != db_owner_name.lower()
                ):

                    warnings.append(
                        f"Row {index}: Owner name differs from system records ({flat_no})"
                    )

                if (
                    phone
                    and db_phone
                    and phone != db_phone
                ):

                    warnings.append(
                        f"Row {index}: Phone differs from system records ({flat_no})"
                    )

            maintenance_amount = normalize_amount(
                maintenance_raw
            )

            interest_amount = normalize_amount(
                interest_raw
            )

            penalty_amount = normalize_amount(
                penalty_raw
            )

            advance_amount = normalize_amount(
                advance_raw
            )

            parsed_rows.append({
                "flat": flat,
                "maintenance_amount": maintenance_amount,
                "interest_amount": interest_amount,
                "penalty_amount": penalty_amount,
                "advance_amount": advance_amount,
            })

        except ValueError as e:

            errors.append(
                f"Row {index}: {str(e)}"
            )

        except Exception as e:

            errors.append(
                f"Row {index}: {str(e)}"
            )

    return {
        "status": "success" if not errors else "failed",
        "errors": errors,
        "warnings": warnings,
        "rows": parsed_rows,
    }

@transaction.atomic
def persist_maintenance_bill_opening_balances(
    society,
    parsed_rows,
):

    from society.models import MemberReceivable

    # ---------------------------------------------
    # DELETE PREVIOUS OPENING BALANCE ENTRIES
    # ---------------------------------------------

    MemberReceivable.objects.filter(
        society=society,
        source_type__startswith="OPENING_",
    ).delete()

    receivables_created = 0

    initialization_rows = 0

    total_outstanding = Decimal("0.00")

    total_advance = Decimal("0.00")

    for row in parsed_rows:

        flat = row["flat"]

        maintenance_amount = row["maintenance_amount"]

        interest_amount = row["interest_amount"]

        penalty_amount = row["penalty_amount"]

        advance_amount = row["advance_amount"]

        # ---------------------------------------------
        # ZERO INITIALIZATION STATE
        # ---------------------------------------------

        if (
            maintenance_amount == 0
            and interest_amount == 0
            and penalty_amount == 0
            and advance_amount == 0
        ):

            MemberReceivable.objects.create(
                society=society,
                flat=flat,
                bill=None,
                source_type="OPENING_INITIALIZATION",
                source_reference="SCR31_OPENING_UPLOAD",
                amount=Decimal("0.00"),
                outstanding_amount=Decimal("0.00"),
                status="PAID",
            )

            initialization_rows += 1

            receivables_created += 1

            continue

        # ---------------------------------------------
        # MAINTENANCE OUTSTANDING
        # ---------------------------------------------

        if maintenance_amount > 0:

            MemberReceivable.objects.create(
                society=society,
                flat=flat,
                bill=None,
                source_type="OPENING_MAINTENANCE",
                source_reference="SCR31_OPENING_UPLOAD",
                amount=maintenance_amount,
                outstanding_amount=maintenance_amount,
                status="OPEN",
            )

            receivables_created += 1

            total_outstanding += maintenance_amount

        # ---------------------------------------------
        # INTEREST OUTSTANDING
        # ---------------------------------------------

        if interest_amount > 0:

            MemberReceivable.objects.create(
                society=society,
                flat=flat,
                bill=None,
                source_type="OPENING_INTEREST",
                source_reference="SCR31_OPENING_UPLOAD",
                amount=interest_amount,
                outstanding_amount=interest_amount,
                status="OPEN",
            )

            receivables_created += 1

            total_outstanding += interest_amount

        # ---------------------------------------------
        # PENALTY OUTSTANDING
        # ---------------------------------------------

        if penalty_amount > 0:

            MemberReceivable.objects.create(
                society=society,
                flat=flat,
                bill=None,
                source_type="OPENING_PENALTY",
                source_reference="SCR31_OPENING_UPLOAD",
                amount=penalty_amount,
                outstanding_amount=penalty_amount,
                status="OPEN",
            )

            receivables_created += 1

            total_outstanding += penalty_amount

        # ---------------------------------------------
        # ADVANCE BALANCE
        # ---------------------------------------------

        if advance_amount > 0:

            advance_value = advance_amount * Decimal("-1")

            MemberReceivable.objects.create(
                society=society,
                flat=flat,
                bill=None,
                source_type="OPENING_ADVANCE",
                source_reference="SCR31_OPENING_UPLOAD",
                amount=advance_value,
                outstanding_amount=advance_value,
                status="OPEN",
            )

            receivables_created += 1

            total_advance += advance_amount

    return {
        "status": "success",
        "receivables_created": receivables_created,
        "initialization_rows": initialization_rows,
        "total_outstanding": total_outstanding,
        "total_advance": total_advance,
    }